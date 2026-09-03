"""Excel and CSV exporters for EMA regulatory document indexes."""

from __future__ import annotations

import csv
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ema_downloader.config import AppConfig
from ema_downloader.database import Database
from ema_downloader.models import DownloadStatus, EMADocument, SyncSummary

logger = logging.getLogger(__name__)

EXCEL_HEADERS = [
    "EMA ID",
    "Document Name",
    "Type",
    "Status",
    "Reference Number",
    "First Published Date",
    "Last Updated Date",
    "Official URL",
    "Local File",
    "Category",
    "Download Status",
    "File Size (Bytes)",
    "SHA-256",
    "Local Hyperlink",
]


def doc_to_export_row(doc: EMADocument, library_dir: Path) -> Dict[str, Any]:
    """Convert an EMADocument to a flat dictionary for export."""
    local_file_path = ""
    local_hyperlink = ""
    if doc.local_path:
        full_path = library_dir / doc.local_path
        local_file_path = str(doc.local_path).replace("\\", "/")
        if full_path.exists():
            # Create a relative link formula or path
            # Excel HYPERLINK formula format: =HYPERLINK("relative_or_abs_path", "Open File")
            # In index folder (00_Index), relative path goes up one level: ../
            rel_from_index = f"../{local_file_path}"
            local_hyperlink = f'=HYPERLINK("{rel_from_index}", "Open File")'

    return {
        "EMA ID": doc.ema_id,
        "Document Name": doc.name,
        "Type": doc.document_type,
        "Status": doc.status,
        "Reference Number": doc.reference_number,
        "First Published Date": doc.first_published_at or "",
        "Last Updated Date": doc.last_updated_at or "",
        "Official URL": doc.official_url,
        "Local File": local_file_path,
        "Category": doc.category or doc.sub_category,
        "Download Status": doc.download_status,
        "File Size (Bytes)": doc.file_size,
        "SHA-256": doc.sha256,
        "Local Hyperlink": local_hyperlink,
    }


class Exporter:
    """Exports document indexes to formatted Excel and CSV files."""

    def __init__(self, config: AppConfig, database: Database) -> None:
        self.config = config
        self.db = database
        self.library_dir = config.storage.library_dir

    def export_csv(self, output_path: Optional[Path] = None) -> Path:
        """Export all documents to CSV."""
        path = output_path or self.config.storage.csv_path
        path.parent.mkdir(parents=True, exist_ok=True)

        docs = self.db.get_documents()
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=EXCEL_HEADERS)
            writer.writeheader()
            for doc in docs:
                row = doc_to_export_row(doc, self.library_dir)
                writer.writerow(row)

        logger.info("Exported %d documents to CSV: %s", len(docs), path)
        return path

    def export_failures_csv(self, output_path: Optional[Path] = None) -> Path:
        """Export failed or error downloads to CSV."""
        path = output_path or self.config.storage.failures_csv_path
        path.parent.mkdir(parents=True, exist_ok=True)

        failure_statuses = [
            DownloadStatus.HTTP_ERROR.value,
            DownloadStatus.INVALID_CONTENT.value,
            DownloadStatus.CHECKSUM_ERROR.value,
            DownloadStatus.WRITE_ERROR.value,
            DownloadStatus.NO_DOWNLOAD_URL.value,
        ]
        failed_docs = self.db.get_documents(download_statuses=failure_statuses)

        fieldnames = ["EMA ID", "Document Name", "Type", "Download Status", "Official URL", "Error Message"]
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for doc in failed_docs:
                writer.writerow({
                    "EMA ID": doc.ema_id,
                    "Document Name": doc.name,
                    "Type": doc.document_type,
                    "Download Status": doc.download_status,
                    "Official URL": doc.official_url,
                    "Error Message": doc.error_message,
                })

        logger.info("Exported %d failure records to CSV: %s", len(failed_docs), path)
        return path

    def export_manifest_csv(self, output_path: Optional[Path] = None) -> Path:
        """Export complete download manifest to CSV."""
        path = output_path or self.config.storage.manifest_csv_path
        path.parent.mkdir(parents=True, exist_ok=True)

        docs = self.db.get_documents()
        fieldnames = [
            "EMA ID", "Reference Number", "Title", "Local Path",
            "File Size", "SHA-256", "Status", "Downloaded At"
        ]
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for doc in docs:
                is_downloaded = doc.download_status in (DownloadStatus.DOWNLOADED.value, DownloadStatus.UPDATED.value)
                has_valid_file = bool(doc.local_path and (self.library_dir / doc.local_path).exists())
                if is_downloaded or has_valid_file:
                    writer.writerow({
                        "EMA ID": doc.ema_id,
                        "Reference Number": doc.reference_number,
                        "Title": doc.name,
                        "Local Path": doc.local_path,
                        "File Size": doc.file_size,
                        "SHA-256": doc.sha256,
                        "Status": doc.download_status,
                        "Downloaded At": doc.downloaded_at or "",
                    })

        logger.info("Exported download manifest to CSV: %s", path)
        return path

    def export_excel(
        self,
        output_path: Optional[Path] = None,
        summary: Optional[SyncSummary] = None,
    ) -> Path:
        """Generate multi-sheet styled Excel index."""
        path = output_path or self.config.storage.excel_path
        path.parent.mkdir(parents=True, exist_ok=True)

        all_docs = self.db.get_documents()
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Style definitions
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        cell_font = Font(name="Segoe UI", size=10)
        link_font = Font(name="Segoe UI", size=10, color="0563C1", underline="single")
        border_side = Side(border_style="thin", color="D9D9D9")
        thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

        def populate_sheet(ws: Any, docs: List[EMADocument], title: str) -> None:
            ws.title = title
            ws.views.sheetView[0].showGridLines = True
            ws.freeze_panes = "A2"

            # Write header
            for col_num, h in enumerate(EXCEL_HEADERS, start=1):
                cell = ws.cell(row=1, column=col_num, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            # Write rows
            for row_num, doc in enumerate(docs, start=2):
                data = doc_to_export_row(doc, self.library_dir)
                for col_num, h in enumerate(EXCEL_HEADERS, start=1):
                    val = data[h]
                    cell = ws.cell(row=row_num, column=col_num)
                    cell.font = cell_font
                    cell.border = thin_border

                    # Handle formulas and hyperlinks
                    if h == "Local Hyperlink" and isinstance(val, str) and val.startswith("="):
                        cell.value = val
                        cell.font = link_font
                    elif h == "Official URL" and val:
                        cell.value = val
                        if val.startswith("http"):
                            cell.hyperlink = val
                            cell.font = link_font
                    else:
                        cell.value = val

            # Adjust column widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or "")
                    if cell.data_type == "f":  # Formula
                        val_str = "Open File"
                    max_len = max(max_len, len(val_str))
                ws.column_dimensions[col_letter].width = max(min(max_len + 4, 45), 12)

        # 1. All Documents
        ws_all = wb.create_sheet(title="All Documents")
        populate_sheet(ws_all, all_docs, "All Documents")

        # 2. Scientific Guidelines
        sci_docs = [d for d in all_docs if d.document_type == "scientific-guideline"]
        ws_sci = wb.create_sheet(title="Scientific Guidelines")
        populate_sheet(ws_sci, sci_docs, "Scientific Guidelines")

        # 3. Regulatory Procedural
        reg_docs = [d for d in all_docs if d.document_type == "regulatory-procedural-guideline"]
        ws_reg = wb.create_sheet(title="Regulatory Procedural")
        populate_sheet(ws_reg, reg_docs, "Regulatory Procedural")

        # 4. Draft and Consultation
        draft_docs = [
            d for d in all_docs
            if "draft" in (d.status or "").lower() or "consultation" in (d.status or "").lower()
        ]
        ws_draft = wb.create_sheet(title="Draft and Consultation")
        populate_sheet(ws_draft, draft_docs, "Draft and Consultation")

        # 5. Failed Downloads
        failure_statuses = {
            DownloadStatus.HTTP_ERROR.value,
            DownloadStatus.INVALID_CONTENT.value,
            DownloadStatus.CHECKSUM_ERROR.value,
            DownloadStatus.WRITE_ERROR.value,
            DownloadStatus.NO_DOWNLOAD_URL.value,
        }
        failed_docs = [d for d in all_docs if d.download_status in failure_statuses]
        ws_fail = wb.create_sheet(title="Failed Downloads")
        populate_sheet(ws_fail, failed_docs, "Failed Downloads")

        # 6. Sync Summary
        ws_sum = wb.create_sheet(title="Sync Summary")
        ws_sum.views.sheetView[0].showGridLines = True
        summary_title_font = Font(name="Segoe UI", size=14, bold=True, color="1F4E78")
        ws_sum.cell(row=1, column=1, value="EMA Document Sync Summary").font = summary_title_font

        counts = self.db.get_summary_counts()
        summary_rows = [
            ("Total Indexed Documents", counts.get("total", 0)),
            ("Successfully Downloaded", counts.get(f"status_{DownloadStatus.DOWNLOADED.value}", 0)),
            ("Updated Documents", counts.get(f"status_{DownloadStatus.UPDATED.value}", 0)),
            ("Skipped (Unchanged)", counts.get(f"status_{DownloadStatus.SKIPPED_UNCHANGED.value}", 0)),
            ("HTTP Errors", counts.get(f"status_{DownloadStatus.HTTP_ERROR.value}", 0)),
            ("Content / Header Errors", counts.get(f"status_{DownloadStatus.INVALID_CONTENT.value}", 0)),
            ("Scientific Guidelines", counts.get("type_scientific-guideline", 0)),
            ("Regulatory Procedural Guidelines", counts.get("type_regulatory-procedural-guideline", 0)),
        ]

        if summary:
            summary_rows.insert(0, ("Last Sync Timestamp", summary.timestamp))
            summary_rows.append(("Sync Duration (seconds)", f"{summary.duration_seconds:.2f}"))
            summary_rows.append(("Downloaded in Last Run", summary.downloaded_success))
            summary_rows.append(("Failed in Last Run", summary.downloaded_failed))

        for idx, (label, val) in enumerate(summary_rows, start=3):
            cell_k = ws_sum.cell(row=idx, column=1, value=label)
            cell_v = ws_sum.cell(row=idx, column=2, value=val)
            cell_k.font = Font(name="Segoe UI", size=10, bold=True)
            cell_v.font = Font(name="Segoe UI", size=10)

        ws_sum.column_dimensions["A"].width = 32
        ws_sum.column_dimensions["B"].width = 25

        wb.save(str(path))
        logger.info("Saved formatted Excel index to %s", path)
        return path

    def export_all(self, summary: Optional[SyncSummary] = None) -> Dict[str, Path]:
        """Export Excel, CSV, manifest, and failures files."""
        return {
            "excel": self.export_excel(summary=summary),
            "csv": self.export_csv(),
            "failures": self.export_failures_csv(),
            "manifest": self.export_manifest_csv(),
        }

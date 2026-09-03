"""Tests for Excel and CSV export generation."""

import openpyxl
from ema_downloader.config import AppConfig
from ema_downloader.database import Database
from ema_downloader.exporter import Exporter
from ema_downloader.models import DownloadStatus, EMADocument, SyncSummary


def test_export_all_formats(sample_config: AppConfig, sample_db: Database):
    docs = [
        EMADocument(
            ema_id="1",
            name="Scientific Doc 1",
            document_type="scientific-guideline",
            status="Adopted",
            reference_number="REF-1",
            official_url="https://ema.europa.eu/1.pdf",
            local_path="01_Scientific_Guidelines/01_Quality/1.pdf",
            download_status=DownloadStatus.DOWNLOADED.value,
            file_size=2048,
            sha256="hash123",
        ),
        EMADocument(
            ema_id="2",
            name="Regulatory Doc 2",
            document_type="regulatory-procedural-guideline",
            status="Draft",
            reference_number="REF-2",
            official_url="https://ema.europa.eu/2.pdf",
            download_status=DownloadStatus.HTTP_ERROR.value,
            error_message="HTTP 404 Not Found",
        ),
    ]
    sample_db.batch_upsert_documents(docs)

    exporter = Exporter(sample_config, sample_db)
    summary = SyncSummary(
        timestamp="2026-09-03T10:00:00Z",
        total_records_in_feed=10,
        total_filtered=2,
        downloaded_success=1,
        downloaded_failed=1,
        duration_seconds=1.5,
    )

    results = exporter.export_all(summary=summary)

    # Check that all expected output files exist
    assert results["excel"].exists()
    assert results["csv"].exists()
    assert results["failures"].exists()
    assert results["manifest"].exists()

    # Inspect Excel sheets
    wb = openpyxl.load_workbook(str(results["excel"]))
    sheet_names = wb.sheetnames
    assert "All Documents" in sheet_names
    assert "Scientific Guidelines" in sheet_names
    assert "Regulatory Procedural" in sheet_names
    assert "Draft and Consultation" in sheet_names
    assert "Failed Downloads" in sheet_names
    assert "Sync Summary" in sheet_names

    # Check rows in All Documents
    ws_all = wb["All Documents"]
    assert ws_all.max_row == 3  # 1 header + 2 data rows

    # Check Failed Downloads
    ws_fail = wb["Failed Downloads"]
    assert ws_fail.max_row == 2  # 1 header + 1 failure
